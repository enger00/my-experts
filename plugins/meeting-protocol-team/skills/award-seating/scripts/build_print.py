#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
桌牌打印单：由三表工作簿的「座位排布表」生成，供 Word 打印模版邮件合并（MERGEFIELD 姓名）

用法：
  python build_print.py --src 三表.xlsx --template 打印单模板.xlsx --out 打印单.xlsx
可选：
  --sheet 座位排布表      源工作表名
  --first-seat-row 4      从第几排开始取领奖人（与生成时的礼堂留空规则一致）
  --zone 领奖人员区        台下分区名称，用于自动定位领奖区列范围

注意：套用旧模板时必须重置 sheetView.topLeftCell，否则 WPS 打开会跳到空白区域。
"""
import argparse
import re

import openpyxl


def pad(name):
    """两字名中间补两个半角空格，与三字名等宽"""
    n = (name or '').replace(' ', '')
    if name == '待定' or len(n) != 2:
        return name
    return n[0] + '  ' + n[1]


def find_zone(ws, label):
    """按分区标题（如“领奖人员区”）定位领奖区列范围"""
    for row in ws.iter_rows(min_row=1, max_row=15):
        for c in row:
            if isinstance(c.value, str) and label in c.value:
                for mr in ws.merged_cells.ranges:
                    if mr.min_row == c.row and mr.min_col == c.column:
                        return mr.min_col, mr.max_col
                return c.column, c.column
    return 6, 11


def find_seat_rows(ws, first_seat_row):
    """按 A 列“第N排”标记定位排号与起始行"""
    out = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        m = re.match(r'^第(\d+)排$', str(v).strip()) if isinstance(v, str) else None
        if m and int(m.group(1)) >= first_seat_row:
            out.append((int(m.group(1)), r))
    return out


def find_leader_rows(ws):
    """定位主席台“座次号”行与“姓名”行（连续多个整数的一行）"""
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, 20)]
        if sum(1 for v in vals if isinstance(v, int)) >= 5:
            cols = [c for c in range(1, 20) if isinstance(ws.cell(r, c).value, int)]
            return r, r + 1, min(cols), max(cols)
    return 5, 6, 3, 15


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', required=True, help='三表工作簿（座位排布表）')
    ap.add_argument('--template', required=True, help='打印单模板 xlsx（沿用列宽/表名）')
    ap.add_argument('--out', required=True)
    ap.add_argument('--sheet', default='座位排布表')
    ap.add_argument('--first-seat-row', type=int, default=4)
    ap.add_argument('--zone', default='领奖人员区')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.src)
    ws = wb[a.sheet]
    c1, c2 = find_zone(ws, a.zone)
    rows = find_seat_rows(ws, a.first_seat_row)
    lr_no, lr_name, lc1, lc2 = find_leader_rows(ws)
    if not rows:
        raise SystemExit(f'未在第{a.first_seat_row}排及以后找到领奖座位行，请检查 --first-seat-row')

    people = []
    for row_no, r in rows:
        for j, col in enumerate(range(c1, c2 + 1)):
            name = ws.cell(r, col).value
            unit = ws.cell(r + 1, col).value
            if name in (None, ''):
                continue
            people.append({'seat': f'第{row_no}排{j + 1}号', 'unit': unit, 'name': name})

    leaders = []
    for col in range(lc1, lc2 + 1):
        no, nm = ws.cell(lr_no, col).value, ws.cell(lr_name, col).value
        if isinstance(no, int) and nm not in (None, ''):
            leaders.append((no, nm))
    print(f'领奖区列 {c1}-{c2}；座位行 {len(rows)} 排；'
          f'领奖人 {len(people)} 位；主席台领导 {len(leaders)} 位')

    wb2 = openpyxl.load_workbook(a.template)
    s1 = wb2[wb2.sheetnames[0]]
    thin = openpyxl.styles.Side(style='thin', color='BFBFBF')
    BD = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
    F = openpyxl.styles.Font(name='宋体', size=12)
    F_RED = openpyxl.styles.Font(name='宋体', size=12, bold=True, color='FF0000')
    CEN = openpyxl.styles.Alignment(horizontal='center', vertical='center')
    LEFT = openpyxl.styles.Alignment(horizontal='left', vertical='center')
    HDR = openpyxl.styles.Font(name='宋体', size=12, bold=True)
    FILL_HDR = openpyxl.styles.PatternFill('solid', fgColor='D9E2F3')

    for r in range(2, s1.max_row + 1):
        for c in range(1, 16):
            s1.cell(r, c).value = None

    for c, title in [(1, '名单（科室）'), (2, '姓名'), (3, '座位'), (4, '领导（主席台）')]:
        cell = s1.cell(1, c, title)
        cell.font, cell.alignment, cell.fill, cell.border = HDR, CEN, FILL_HDR, BD
    s1.row_dimensions[1].height = 22

    r = 2
    for p in people:
        s1.cell(r, 1, p['unit']).alignment = LEFT
        s1.cell(r, 1).font = F
        cc = s1.cell(r, 2, pad(p['name']))
        cc.alignment = CEN
        cc.font = F_RED if p['name'] == '待定' else F
        s1.cell(r, 3, p['seat']).alignment = CEN
        s1.cell(r, 3).font = F
        for c in range(1, 4):
            s1.cell(r, c).border = BD
        s1.row_dimensions[r].height = 20
        r += 1

    r2 = 2
    for no, nm in sorted(leaders):
        cc = s1.cell(r2, 4, f'{no}. {pad(nm)}')
        cc.font, cc.alignment, cc.border = F, LEFT, BD
        r2 += 1

    s1.column_dimensions['A'].width = 20.75
    s1.column_dimensions['B'].width = 13
    s1.column_dimensions['C'].width = 12
    s1.column_dimensions['D'].width = 15
    s1.freeze_panes = 'A2'

    # Sheet2：领导桌牌（单列“姓名”，同一 Word 模版合并即可）
    if len(wb2.sheetnames) > 1:
        s2 = wb2[wb2.sheetnames[1]]
        for row in s2.iter_rows():
            for c in row:
                c.value = None
        s2.cell(1, 1, '姓名')
        s2.cell(1, 1).font, s2.cell(1, 1).alignment = HDR, CEN
        s2.cell(1, 1).fill, s2.cell(1, 1).border = FILL_HDR, BD
        rr = 2
        for no, nm in sorted(leaders):
            cc = s2.cell(rr, 1, pad(nm))
            cc.font, cc.alignment, cc.border = F, CEN, BD
            rr += 1
        s2.column_dimensions['A'].width = 15

    # 关键：重置视图，避免继承模板的滚动位置导致打开显示空白
    for w in wb2.worksheets:
        w.sheet_view.topLeftCell = 'A1'
        w.sheet_view.zoomScale = 100
    wb2.active = 0
    wb2.save(a.out)
    print('saved:', a.out)
    print('待定:', [p['unit'] for p in people if p['name'] == '待定'] or '无')


if __name__ == '__main__':
    main()
